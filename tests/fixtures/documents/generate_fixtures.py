"""Regenerates the binary document fixtures in this directory.

Run with ``python tests/fixtures/documents/generate_fixtures.py`` from a
venv with the ``dev`` extras installed (python-docx/python-pptx/openpyxl --
all dev/test-only, see pyproject.toml). Not imported by the test suite
itself; the fixtures it writes are committed so tests don't depend on
these libraries' exact output being stable across versions.

Search Quality Improvement Plan, Phase 10: ``_make_odt``/``_make_ods``/
``_make_odp`` use ``odfdo`` -- unlike python-docx/python-pptx/openpyxl
above, this is a genuine *runtime* dependency now (Docling's own ODT/ODS/
ODP backend needs it to convert a project's real files, not just to
generate these fixtures -- see ``documents/docling_adapter.py``'s module
docstring and ``pyproject.toml``'s dependency comment), so it's already
installed in every dev venv rather than needing a separate dev-only
entry. ``_make_epub`` needs nothing beyond the standard library
(``zipfile``): an EPUB is just a ZIP of XHTML, so it's hand-assembled the
same deliberate way ``_make_pdf``/``_make_scanned_pdf`` below are.
``_make_ocr_image`` needs Pillow, already a transitive dependency via
Docling's own image backend.
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).parent


def _make_docx() -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("Doc Title", level=1)
    doc.add_paragraph("Intro paragraph text.")
    doc.add_heading("Section One", level=2)
    doc.add_paragraph("Paragraph in section one.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1"
    doc.save(HERE / "document.docx")


def _make_pptx() -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Presentation Title"
    body = slide.placeholders[1].text_frame
    body.text = "First bullet point"
    body.add_paragraph().text = "Second bullet point"
    prs.save(HERE / "presentation.pptx")


def _make_xlsx() -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["B1"] = "Name", "Value"
    ws["A2"], ws["B2"] = "alpha", 1
    ws["A3"], ws["B3"] = "beta", 2
    wb.save(HERE / "spreadsheet.xlsx")


def _make_pdf() -> None:
    """A minimal, hand-assembled single-page PDF (no reportlab/fpdf
    dependency) with two lines of real text content -- enough for
    docling_pdf-marked golden tests.
    """
    lines = ["Sample PDF Title", "This is a line of body text in the sample PDF."]
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        "/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    y = 750
    content_lines = ["BT", "/F1 18 Tf"]
    for line in lines:
        content_lines.append(f"1 0 0 1 72 {y} Tm ({line}) Tj")
        y -= 24
    content_lines.append("ET")
    content = "\n".join(content_lines)
    objects.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")

    pdf = "%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n{obj}\nendobj\n"
    xref_offset = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    for off in offsets[1:]:
        pdf += f"{off:010} 00000 n \n"
    pdf += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF"
    (HERE / "sample.pdf").write_bytes(pdf.encode("latin-1"))


def _make_scanned_pdf() -> None:
    """A minimal, hand-assembled single-page PDF with a completely empty
    content stream -- no text, no image, nothing Docling's plain pipeline
    (OCR off) could ever extract a character from. Stands in for a real
    scanned/image-only PDF for Search Quality Improvement Plan Phase 5's
    "auto" OCR-trigger golden test (``test_docling_pdf.py``): it isn't a
    raster image of text, but it exercises the exact condition that
    actually triggers OCR -- zero extracted text on every page -- without
    needing a real scanner output or an image-embedding library this
    project doesn't otherwise depend on.
    """
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /Resources << >> "
        "/MediaBox [0 0 612 792] /Contents 4 0 R >>",
        "<< /Length 0 >>\nstream\n\nendstream",
    ]

    pdf = "%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n{obj}\nendobj\n"
    xref_offset = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    for off in offsets[1:]:
        pdf += f"{off:010} 00000 n \n"
    pdf += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF"
    (HERE / "scanned.pdf").write_bytes(pdf.encode("latin-1"))


def _make_odt() -> None:
    from odfdo import Document, Header, Paragraph

    doc = Document("text")
    doc.body.append(Header(1, "Doc Title"))
    doc.body.append(Paragraph("Intro paragraph text."))
    doc.body.append(Header(2, "Section One"))
    doc.body.append(Paragraph("Paragraph in section one."))
    doc.save(HERE / "document.odt")


def _make_ods() -> None:
    from odfdo import Document, Table

    doc = Document("spreadsheet")
    table = Table("Sheet1")
    table.set_values([["Name", "Value"], ["alpha", 1], ["beta", 2]])
    doc.body.append(table)
    doc.save(HERE / "spreadsheet.ods")


def _make_odp() -> None:
    from odfdo import Document, DrawPage, Frame

    doc = Document("presentation")
    page = DrawPage(name="Slide 1")
    page.append(
        Frame.text_frame(
            "Presentation Title", presentation_class="title", size=("20cm", "2cm")
        )
    )
    page.append(
        Frame.text_frame(
            "First bullet point", presentation_class="outline", size=("20cm", "10cm")
        )
    )
    doc.body.append(page)
    doc.save(HERE / "presentation.odp")


def _make_epub() -> None:
    """Hand-assembled the same deliberate way ``_make_pdf`` is: an EPUB is
    just a ZIP archive of a ``container.xml`` pointer, an OPF manifest/
    spine, and XHTML content files -- no ``ebooklib``-style dependency
    needed to build one, matching what Docling's own
    ``EpubDocumentBackend`` (stdlib ``zipfile`` + ``defusedxml``) expects
    to read.
    """
    import zipfile

    container_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<container version="1.0" '
        'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
        "  <rootfiles>\n"
        '    <rootfile full-path="content.opf" '
        'media-type="application/oebps-package+xml"/>\n'
        "  </rootfiles>\n"
        "</container>\n"
    )
    content_opf = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" '
        'unique-identifier="bookid" version="2.0">\n'
        '  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        "    <dc:title>Sample Book</dc:title>\n"
        '    <dc:identifier id="bookid">urn:uuid:sample-book-0001</dc:identifier>\n'
        "    <dc:language>en</dc:language>\n"
        "  </metadata>\n"
        "  <manifest>\n"
        '    <item id="chapter1" href="chapter1.xhtml" '
        'media-type="application/xhtml+xml"/>\n'
        '    <item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>\n'
        "  </manifest>\n"
        '  <spine toc="ncx">\n'
        '    <itemref idref="chapter1"/>\n'
        "  </spine>\n"
        "</package>\n"
    )
    toc_ncx = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">\n'
        '  <head><meta name="dtb:uid" content="urn:uuid:sample-book-0001"/></head>\n'
        "  <docTitle><text>Sample Book</text></docTitle>\n"
        "  <navMap>\n"
        '    <navPoint id="navpoint-1" playOrder="1">\n'
        "      <navLabel><text>Chapter One</text></navLabel>\n"
        '      <content src="chapter1.xhtml"/>\n'
        "    </navPoint>\n"
        "  </navMap>\n"
        "</ncx>\n"
    )
    chapter1 = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<html xmlns="http://www.w3.org/1999/xhtml">\n'
        "<head><title>Chapter One</title></head>\n"
        "<body>\n"
        "<h1>Chapter One</h1>\n"
        "<p>This is the first paragraph of the sample EPUB book.</p>\n"
        "</body>\n"
        "</html>\n"
    )
    with zipfile.ZipFile(HERE / "sample.epub", "w") as zf:
        zf.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("META-INF/container.xml", container_xml)
        zf.writestr("content.opf", content_opf)
        zf.writestr("toc.ncx", toc_ncx)
        zf.writestr("chapter1.xhtml", chapter1)


def _make_ocr_image() -> None:
    """A small PNG with real, rendered (not embedded-as-text) words --
    the only way to prove OCR extraction genuinely reads pixels rather
    than round-tripping empty text (see ``documents/docling_adapter.py``'s
    module docstring on why images need OCR at all). Falls back to
    Pillow's bundled bitmap font when no system TrueType font is found,
    since a CI runner's available fonts aren't guaranteed -- RapidOCR
    reads either just fine at this size/weight.
    """
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (600, 200), color="white")
    draw = ImageDraw.Draw(image)
    font: ImageFont.ImageFont | ImageFont.FreeTypeFont
    try:
        font = ImageFont.truetype(
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 28
        )
    except OSError:
        font = ImageFont.load_default()
    draw.text((20, 30), "Sample Image Title", fill="black", font=font)
    draw.text((20, 90), "OCR body text line here.", fill="black", font=font)
    image.save(HERE / "sample_ocr.png")


if __name__ == "__main__":
    _make_docx()
    _make_pptx()
    _make_xlsx()
    _make_pdf()
    _make_scanned_pdf()
    _make_odt()
    _make_ods()
    _make_odp()
    _make_epub()
    _make_ocr_image()
    print("Fixtures written to", HERE)
